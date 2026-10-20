from datetime import UTC, datetime, timedelta

from sqlalchemy import select

from app.core.config import get_settings
from app.models import AckStatus, CustomerEmail, EmailStatus, Priority
from app.services.app_settings import get_app_settings
from app.services.datasources import DataSources
from app.services.processor import poll_once
from tests.conftest import BALANCE_TOPIC, FAQ_TOPIC


def _poll(db, mailbox):
    return poll_once(db, mailbox, DataSources(get_settings()), get_settings())


def _start(db, mailbox, *, auto_reply=True):
    """First poll only sets the starting point; then switch automatic answers on or off."""
    _poll(db, mailbox)
    cfg = get_app_settings(db)
    cfg.auto_reply_enabled = auto_reply
    db.commit()


def _email(db, msg_id) -> CustomerEmail:
    return db.scalar(select(CustomerEmail).where(CustomerEmail.graph_message_id == msg_id))


def test_first_poll_does_not_touch_old_mail(db, mailbox):
    mailbox.add(
        "old",
        subject="Old",
        body="Sent before the system started",
        received=datetime.now(UTC) - timedelta(days=3),
    )
    assert _poll(db, mailbox) == 0
    assert mailbox.sent == []
    assert _email(db, "old") is None


def test_general_question_is_acknowledged_and_answered_from_the_database(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.add("m1", subject="Opening hours", body="Hello,\nWhat time do your branches open?\nThanks")
    assert _poll(db, mailbox) == 1

    email = _email(db, "m1")
    assert email.ack_status == AckStatus.SENT
    assert email.status == EmailStatus.AUTO_ANSWERED
    assert len(mailbox.sent) == 2
    ack_html, answer_html = mailbox.sent[0][1], mailbox.sent[1][1]
    assert email.reference in ack_html
    assert "08:30 to 16:00" in answer_html  # straight from the bank database


def test_customer_specific_answer_uses_values_from_the_email(db, mailbox, live_topic):
    live_topic(BALANCE_TOPIC)
    _start(db, mailbox)
    mailbox.add("m2", subject="Balance", body="Please what is my account balance? Account 6200123456.")
    _poll(db, mailbox)
    assert _email(db, "m2").status == EmailStatus.AUTO_ANSWERED
    assert "1,520.50" in mailbox.sent[-1][1]
    assert "Savings" in mailbox.sent[-1][1]


def test_customer_data_only_goes_to_the_registered_email(db, mailbox, live_topic):
    live_topic(BALANCE_TOPIC)
    _start(db, mailbox)
    # Kofi asks about Ama's account: the query also matches on the sender's email, so no data.
    mailbox.add(
        "m3",
        subject="Balance",
        body="What is my account balance? 6200123456",
        sender="kofi@example.com",
        name="Kofi",
    )
    _poll(db, mailbox)
    email = _email(db, "m3")
    assert email.status == EmailStatus.QUEUED
    assert "No matching record" in email.status_reason
    assert len(mailbox.sent) == 1  # acknowledgement only


def test_unverified_sender_never_receives_customer_data(db, mailbox, live_topic):
    live_topic(BALANCE_TOPIC)
    _start(db, mailbox)
    mailbox.add(
        "m4",
        subject="Balance",
        body="What is my account balance? 6200123456",
        headers={"Authentication-Results": "spf=fail; dmarc=fail"},
    )
    _poll(db, mailbox)
    email = _email(db, "m4")
    assert email.status == EmailStatus.QUEUED
    assert "could not be verified" in email.status_reason
    assert all("1,520" not in html for _, html in mailbox.sent)


def test_missing_value_in_email_goes_to_officer(db, mailbox, live_topic):
    live_topic(BALANCE_TOPIC)
    _start(db, mailbox)
    mailbox.add("m5", subject="Balance", body="What is my account balance?")
    _poll(db, mailbox)
    email = _email(db, "m5")
    assert email.status == EmailStatus.QUEUED
    assert "account number" in email.status_reason


def test_sensitive_email_is_high_priority_and_never_auto_answered(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.add("m6", subject="Fraud on my card", body="What time do your branches open? Someone stole money")
    _poll(db, mailbox)
    email = _email(db, "m6")
    assert email.priority == Priority.HIGH
    assert email.status == EmailStatus.QUEUED
    assert len(mailbox.sent) == 1


def test_unrelated_email_is_queued_with_the_reason(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.add("m7", subject="Job application", body="I would like to apply for the graduate programme.")
    _poll(db, mailbox)
    email = _email(db, "m7")
    assert email.status == EmailStatus.QUEUED
    assert "Nothing matches closely enough" in email.status_reason


def test_auto_answers_switched_off_still_acknowledges(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox, auto_reply=False)
    mailbox.add("m8", subject="Hours", body="What time do your branches open?")
    _poll(db, mailbox)
    email = _email(db, "m8")
    assert email.status == EmailStatus.QUEUED
    assert email.match_topic_id is not None  # suggested to the officer
    assert len(mailbox.sent) == 1


def test_auto_replies_and_bounces_are_skipped(db, mailbox):
    _start(db, mailbox)
    mailbox.add("a1", subject="Out of office", body="I am away", headers={"Auto-Submitted": "auto-replied"})
    mailbox.add("a2", subject="Undeliverable", body="...", sender="MAILER-DAEMON@example.com")
    mailbox.add("a3", subject="Loop", body="...", sender="care@fnb.test")
    _poll(db, mailbox)
    assert all(_email(db, i).status == EmailStatus.SKIPPED for i in ("a1", "a2", "a3"))
    assert mailbox.sent == []


def test_same_email_is_never_processed_twice(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.add("dup", subject="Hours", body="What time do your branches open?")
    _poll(db, mailbox)
    mailbox.new_ids.append("dup")  # Outlook reports it again
    assert _poll(db, mailbox) == 0
    assert len(mailbox.sent) == 2


def test_outlook_failure_is_recorded_and_queued(db, mailbox, live_topic):
    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.fail_replies = True
    mailbox.add("f1", subject="Hours", body="What time do your branches open?")
    _poll(db, mailbox)
    email = _email(db, "f1")
    assert email.ack_status == AckStatus.FAILED
    assert email.status == EmailStatus.QUEUED


def test_needs_attention_groups_reasons_and_suggests_a_copyable_answer(db, mailbox, live_topic, login):
    topic_id = live_topic(FAQ_TOPIC)
    _start(db, mailbox, auto_reply=False)
    mailbox.add("o1", subject="Hours", body="What time do your branches open?")
    mailbox.add("o2", subject="Hello", body="Can someone call me back about my loan application?")
    _poll(db, mailbox)
    officer = login("olivia")

    att = officer.get("/api/emails/attention").json()
    assert att["page"]["total"] == 2
    assert {r["code"] for r in att["reasons"]} == {"auto_off", "no_match"}
    only = officer.get("/api/emails/attention?reason=no_match").json()
    assert [e["reference"] for e in only["page"]["items"]] == [_email(db, "o2").reference]

    first = _email(db, "o1")
    suggestion = officer.get(f"/api/emails/{first.id}/suggestion/{topic_id}").json()
    assert "08:30" in suggestion["text"]

    resp = officer.post(f"/api/emails/{first.id}/resolve", json={"outcome": "replied", "note": "Called her"})
    assert resp.json()["status"] == "replied" and resp.json()["handled_via"] == "dashboard"
    second = _email(db, "o2")
    assert officer.post(f"/api/emails/{second.id}/resolve", json={"outcome": "closed"}).status_code == 422
    assert (
        officer.post(f"/api/emails/{second.id}/resolve", json={"outcome": "closed", "note": "Spam"}).json()[
            "status"
        ]
        == "closed"
    )
    assert officer.get("/api/emails/attention").json()["page"]["total"] == 0


def test_reply_written_in_outlook_marks_the_email_handled(db, mailbox):
    from app.services import outlook_replies
    from app.services.graph_client import SentItem

    _start(db, mailbox)
    mailbox.add("h1", subject="Loan", body="Please call me about a loan.")
    _poll(db, mailbox)
    email = _email(db, "h1")
    assert email.status == EmailStatus.QUEUED
    ack = email.replies[0].body_text
    later = datetime.now(UTC) + timedelta(minutes=5)

    # Only the assistant's own acknowledgement is in Sent Items: still waiting.
    mailbox.sent_items["conv-h1"] = [SentItem(sent_at=later, preview=ack[:200])]
    assert outlook_replies.check_waiting(db, mailbox) == 0

    # A person replied in Outlook.
    mailbox.sent_items["conv-h1"].append(SentItem(sent_at=later, preview="Hi Ama, I will call you at 2pm."))
    assert outlook_replies.check_waiting(db, mailbox) == 1
    db.refresh(email)
    assert email.status == EmailStatus.REPLIED and email.handled_via == "outlook"


def test_follow_up_in_same_thread_is_not_acknowledged_twice(db, mailbox):
    _start(db, mailbox)
    mailbox.add("t1", subject="Loan", body="Please call me about a loan.")
    _poll(db, mailbox)
    msg = mailbox.messages["t1"]
    mailbox.add("t2", subject="RE: Loan", body="Any update?")
    mailbox.messages["t2"].conversation_id = msg.conversation_id
    _poll(db, mailbox)
    follow = _email(db, "t2")
    assert follow.is_follow_up and follow.ack_status == AckStatus.NOT_APPLICABLE
    assert len(mailbox.sent) == 1  # only the first email was acknowledged


def test_exports_are_pdf_and_excel_with_the_filters_on_screen(db, mailbox, login):
    import io

    from openpyxl import load_workbook

    _start(db, mailbox)
    mailbox.add("x1", subject='=HYPERLINK("http://evil")', body="hi there friend")
    mailbox.add("x2", subject="FRAUD", body="someone stole money")
    _poll(db, mailbox)
    admin = login("adam")

    pdf = admin.get("/api/emails/export.pdf?scope=attention")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "needs-attention" in pdf.headers["content-disposition"]

    xlsx = admin.get("/api/emails/export.xlsx?scope=attention&filter=urgent")
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws["A1"].value.endswith("Smart Email Assistance") and ws["A2"].value == "Needs attention"
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
    assert len(rows) == 1 and rows[0][4] == "FRAUD"  # only the urgent one

    log = load_workbook(io.BytesIO(admin.get("/api/emails/export.xlsx").content)).active
    subjects = [r[3].value for r in log.iter_rows(min_row=6)]
    assert '\'=HYPERLINK("http://evil")' in subjects  # formulas are stored as text


def test_failure_while_processing_hands_the_email_to_a_person(db, mailbox, live_topic, monkeypatch):
    from app.services import processor

    live_topic(FAQ_TOPIC)
    _start(db, mailbox)

    def broken(*args, **kwargs):
        raise RuntimeError("matcher crashed")

    monkeypatch.setattr(processor, "process_email", broken)
    mailbox.add("crash", subject="Hours", body="What time do your branches open?")
    mailbox.add("ok", subject="Hours", body="What time do your branches open?")
    _poll(db, mailbox)
    db.expire_all()
    for msg_id in ("crash", "ok"):  # one failure does not stop the rest of the batch
        email = _email(db, msg_id)
        assert email.status == EmailStatus.QUEUED  # never left invisible as "Processing"
        assert email.reason_code == "interrupted"


def test_email_stored_but_not_processed_is_picked_up_next_cycle(db, mailbox, live_topic):
    from app.services.processor import store_message

    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.add("late", subject="Hours", body="What time do your branches open?")
    mailbox.new_ids.clear()  # stored by a cycle that then failed; Outlook will not report it again
    store_message(db, mailbox.messages["late"])
    assert _poll(db, mailbox) == 1
    db.expire_all()
    assert _email(db, "late").status == EmailStatus.AUTO_ANSWERED
    assert len(mailbox.sent) == 2  # acknowledgement and answer, once each


def test_message_deleted_before_it_is_read_does_not_block_the_mailbox(db, mailbox, live_topic):
    from app.services.graph_client import GraphError

    live_topic(FAQ_TOPIC)
    _start(db, mailbox)
    mailbox.new_ids.append("gone")
    original = mailbox.get_message

    def get_message(message_id):
        if message_id == "gone":
            raise GraphError("not found", status=404)
        return original(message_id)

    mailbox.get_message = get_message
    mailbox.add("next", subject="Hours", body="What time do your branches open?")
    assert _poll(db, mailbox) == 1
    assert _email(db, "next").status == EmailStatus.AUTO_ANSWERED

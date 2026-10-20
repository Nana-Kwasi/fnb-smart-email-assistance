"""PDF and Excel exports with the FNB layout: logo, application name, title, who and when."""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

APP_NAME = "Smart Email Assistance"
LOGO = Path(__file__).resolve().parent.parent / "assets" / "fnb-logo.png"
TEAL = "#00857F"
TEAL_DEEP = "#00615C"
ORANGE = "#F58220"
INK = "#14232B"
MUTED = "#6B7B83"
ZEBRA = "#F2F7F7"


@dataclass
class Column:
    title: str
    width: float  # relative width (PDF) and roughly characters (Excel)


@dataclass
class Report:
    title: str
    subtitle: str  # e.g. the filters applied
    generated_by: str
    generated_at: datetime
    columns: list[Column]
    rows: list[list[str]]


# -- Excel --------------------------------------------------------------------------------------


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    last_col = get_column_letter(len(report.columns))

    ws["A1"] = f"FNB · {APP_NAME}"
    ws["A1"].font = Font(bold=True, size=14, color=TEAL_DEEP[1:])
    ws["A2"] = report.title
    ws["A2"].font = Font(bold=True, size=12, color=INK[1:])
    ws["A3"] = (
        f"{report.subtitle}   ·   Generated {report.generated_at:%d %b %Y, %H:%M} by {report.generated_by}"
    )
    ws["A3"].font = Font(size=10, color=MUTED[1:])
    for r in (1, 2, 3):
        ws.merge_cells(f"A{r}:{last_col}{r}")

    header_row = 5
    thin = Side(style="thin", color="DCE7E5")
    for i, col in enumerate(report.columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=col.title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL[1:])
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(i)].width = max(10, col.width)
    ws.row_dimensions[header_row].height = 22

    for r, row in enumerate(report.rows, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            # A leading = + - @ would make Excel run it as a formula; store it as text instead.
            text = "" if value is None else str(value)
            cell = ws.cell(row=r, column=c, value=("'" + text) if text[:1] in "=+-@" else text)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if (r - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ZEBRA[1:])

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + max(len(report.rows), 1)}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -- PDF ----------------------------------------------------------------------------------------


def to_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    page_w, page_h = landscape(A4)
    margin = 14 * mm
    header_h = 24 * mm
    logo = ImageReader(str(LOGO)) if LOGO.exists() else None

    cell = ParagraphStyle(
        "cell",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor(INK),
        alignment=TA_LEFT,
    )
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    sub = ParagraphStyle("sub", parent=cell, fontSize=9, leading=12, textColor=colors.HexColor(MUTED))

    class NumberedCanvas(pdfcanvas.Canvas):
        """Draws the header and footer once the total page count is known."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):  # noqa: N802 - reportlab API
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total: int) -> None:
            top = page_h - margin
            x = margin
            if logo:
                iw, ih = logo.getSize()
                h = 13 * mm
                w = h * iw / ih
                self.drawImage(logo, x, top - h, width=w, height=h, mask="auto")
                x += w + 6 * mm
                self.setStrokeColor(colors.HexColor("#DCE7E5"))
                self.line(x - 3 * mm, top - h, x - 3 * mm, top)
            self.setFillColor(colors.HexColor(TEAL_DEEP))
            self.setFont("Helvetica-Bold", 13)
            self.drawString(x, top - 5 * mm, APP_NAME)
            self.setFillColor(colors.HexColor(INK))
            self.setFont("Helvetica-Bold", 11)
            self.drawString(x, top - 11 * mm, report.title)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor(MUTED))
            right = page_w - margin
            self.drawRightString(right, top - 5 * mm, f"Generated {report.generated_at:%d %b %Y, %H:%M}")
            self.drawRightString(right, top - 9.5 * mm, f"by {report.generated_by}")
            self.setStrokeColor(colors.HexColor(ORANGE))
            self.setLineWidth(1.5)
            self.line(margin, top - header_h + 6 * mm, right, top - header_h + 6 * mm)
            # Footer
            self.setLineWidth(0.5)
            self.setStrokeColor(colors.HexColor("#DCE7E5"))
            self.line(margin, margin - 2 * mm, right, margin - 2 * mm)
            self.setFont("Helvetica", 7.5)
            self.setFillColor(colors.HexColor(MUTED))
            self.drawString(margin, margin - 6 * mm, "Confidential · FNB internal use only")
            self.drawRightString(right, margin - 6 * mm, f"Page {self._pageNumber} of {total}")

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=(page_w, page_h),
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin + header_h - 4 * mm,
        bottomMargin=margin + 2 * mm,
        title=f"{report.title} - {APP_NAME}",
        author="FNB",
        creator=APP_NAME,
    )
    usable = page_w - 2 * margin
    total_weight = sum(c.width for c in report.columns)
    widths = [usable * c.width / total_weight for c in report.columns]

    def esc(text: str) -> str:
        return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    data = [[Paragraph(esc(c.title), head) for c in report.columns]]
    data += [[Paragraph(esc(str(v)), cell) for v in row] for row in report.rows]
    if not report.rows:
        data.append(
            [Paragraph("Nothing to show for these filters.", cell)] + [""] * (len(report.columns) - 1)
        )
    table = LongTable(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(TEAL)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#DCE7E5")),
    ]
    style += [("BACKGROUND", (0, r), (-1, r), colors.HexColor(ZEBRA)) for r in range(2, len(data), 2)]
    if not report.rows:
        style.append(("SPAN", (0, 1), (-1, 1)))
    table.setStyle(TableStyle(style))
    story = [Paragraph(esc(report.subtitle), sub), Spacer(1, 4 * mm), table]
    doc.build(story, canvasmaker=NumberedCanvas)
    return out.getvalue()
